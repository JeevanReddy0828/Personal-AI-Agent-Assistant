from __future__ import annotations

import asyncio
import html as html_lib
import re
import time
from collections.abc import Callable
from pathlib import Path

from laptop_agent.safety import ApprovalGate, ApprovalRequest, RiskLevel
from laptop_agent.tools.base import ToolResult

# The model writes the document body as Markdown; injected so the success path is
# unit-tested offline, per the weather/websearch/imagegen pattern.
Writer = Callable[[str], str]

FORMATS = {"pdf": "application/pdf",
           "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           "md": "text/markdown; charset=utf-8"}
# "as a word doc", "in pdf", "as markdown" — the format is usually the tail of the request.
_FORMAT_WORDS = {
    "pdf": "pdf", "word": "docx", "docx": "docx", "doc": "docx",
    "markdown": "md", "md": "md", "text": "md",
}
_FORMAT_TAIL = re.compile(
    # ^ as well as \s+, so "document as a pdf" is recognised as having no subject at all
    # rather than being written up as a document about the words "as a pdf".
    r"(?:^|\s+)(?:as|in|to|into)\s+(?:an?\s+)?(pdf|word(?:\s+doc(?:ument)?)?|docx|doc|markdown|md|text)"
    r"(?:\s+(?:file|doc|document|format))?\s*$",
    re.IGNORECASE,
)

_PROMPT = (
    "Write a complete, well-structured document for this request:\n\n{spec}\n\n"
    "Rules:\n"
    "- Output Markdown only. No preamble, no commentary, no code fence around the whole thing.\n"
    "- Start with a single `# Title` line.\n"
    "- Use `##` sections, short paragraphs, `-` bullets and Markdown tables where they earn "
    "their place.\n"
    "- Be specific and complete. Do not leave placeholders like [insert X]."
)

# A print stylesheet, not the app's: this is read on paper, so serif body text, real
# margins and headings that stay with their section.
_CSS = """
@page{size:Letter;margin:0.9in 0.85in}
*{box-sizing:border-box}
body{font:11.5pt/1.55 Georgia,'Times New Roman',serif;color:#14181d;margin:0}
h1{font:700 22pt/1.25 Georgia,serif;margin:0 0 4pt;letter-spacing:-.2pt}
h2{font:700 14pt/1.3 Georgia,serif;margin:20pt 0 6pt;border-bottom:.6pt solid #c8d0d8;padding-bottom:3pt;break-after:avoid}
h3{font:700 12pt/1.3 Georgia,serif;margin:14pt 0 4pt;break-after:avoid}
p{margin:0 0 8pt;orphans:2;widows:2}
ul,ol{margin:0 0 8pt;padding-left:18pt}
li{margin:0 0 3pt}
code{font:10pt ui-monospace,Consolas,monospace;background:#f2f4f7;padding:1pt 3pt;border-radius:3px}
pre{font:9.5pt/1.4 ui-monospace,Consolas,monospace;background:#f5f7f9;border:.6pt solid #dde3e9;
    border-radius:4px;padding:8pt 10pt;margin:0 0 10pt;white-space:pre-wrap;break-inside:avoid}
table{border-collapse:collapse;width:100%;margin:0 0 10pt;font-size:10.5pt;break-inside:avoid}
th,td{border:.6pt solid #c8d0d8;padding:5pt 7pt;text-align:left;vertical-align:top}
th{background:#eef2f5;font-weight:700}
.meta{color:#68727d;font-size:9.5pt;margin:0 0 16pt}
"""


def _slug(text: str, limit: int = 48) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return cleaned[:limit].strip("-") or "document"


def split_format(spec: str, default: str = "pdf") -> tuple[str, str]:
    """Pull a trailing format off the request: ('a report on rust', 'pdf')."""
    text = (spec or "").strip()
    match = _FORMAT_TAIL.search(text)
    if not match:
        return text, default
    word = re.sub(r"\s+doc(ument)?$", "", match.group(1).strip().lower())
    return text[: match.start()].strip(), _FORMAT_WORDS.get(word, default)


def _inline(text: str) -> str:
    out = html_lib.escape(text)
    out = re.sub(r"`([^`]+)`", r"<code>\1</code>", out)
    out = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", out)
    out = re.sub(r"(?<![\w*])\*([^*]+)\*", r"<em>\1</em>", out)
    out = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", out)  # a printed page cannot be clicked
    return out


def markdown_to_html(md: str, title: str = "") -> str:
    """Enough Markdown for a written document: headings, lists, tables, code, rules."""
    lines = (md or "").replace("\r\n", "\n").split("\n")
    body: list[str] = []
    list_tag = None
    index = 0

    def close_list() -> None:
        nonlocal list_tag
        if list_tag:
            body.append(f"</{list_tag}>")
            list_tag = None

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if line.startswith("```"):
            close_list()
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            body.append("<pre>" + html_lib.escape("\n".join(block)) + "</pre>")
            index += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            close_list()
            level = min(len(heading.group(1)), 4)
            body.append(f"<h{level}>{_inline(heading.group(2))}</h{level}>")
            index += 1
            continue
        # | a | b |  over  |---|---|
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|?$", lines[index + 1].strip()):
            close_list()
            cells = lambda row: [c.strip() for c in row.strip().strip("|").split("|")]
            head = cells(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(cells(lines[index]))
                index += 1
            head_html = "".join(f"<th>{_inline(c)}</th>" for c in head)
            rows_html = "".join("<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in r) + "</tr>" for r in rows)
            body.append(f"<table><thead><tr>{head_html}</tr></thead><tbody>{rows_html}</tbody></table>")
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", raw)
        if bullet:
            if list_tag != "ul":
                close_list()
                body.append("<ul>")
                list_tag = "ul"
            body.append(f"<li>{_inline(bullet.group(1))}</li>")
            index += 1
            continue
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", raw)
        if numbered:
            if list_tag != "ol":
                close_list()
                body.append("<ol>")
                list_tag = "ol"
            body.append(f"<li>{_inline(numbered.group(1))}</li>")
            index += 1
            continue
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            close_list()
            index += 1
            continue
        close_list()
        body.append(f"<p>{_inline(line)}</p>")
        index += 1
    close_list()
    head = f"<title>{html_lib.escape(title)}</title>" if title else ""
    return f"<!doctype html><html><head><meta charset='utf-8'>{head}<style>{_CSS}</style></head><body>{''.join(body)}</body></html>"


def _title_of(md: str, fallback: str) -> str:
    for line in (md or "").split("\n"):
        heading = re.match(r"^#\s+(.*)$", line.strip())
        if heading:
            return heading.group(1).strip()
    return fallback


def _write_docx(md: str, out_path: Path) -> ToolResult:
    try:
        from docx import Document  # type: ignore
        from docx.shared import Pt
    except ImportError as exc:
        # The abandoned PyPI package named `docx` shadows python-docx and fails on import
        # under Python 3, so name the real reason instead of only repeating the install line.
        return ToolResult.failure(
            f"Word export needs python-docx and could not import it ({exc}). "
            "Run: pip uninstall docx && pip install python-docx "
            "(the old 'docx' package shadows it). Ask for a PDF instead and it works now."
        )
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    lines = (md or "").replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|?$", lines[index + 1].strip()):
            cells = lambda row: [c.strip() for c in row.strip().strip("|").split("|")]
            head = cells(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(cells(lines[index]))
                index += 1
            table = document.add_table(rows=1, cols=len(head))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, head):
                cell.text = text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            for row in rows:
                target = table.add_row().cells
                for cell, text in zip(target, row):
                    cell.text = text
            document.add_paragraph()
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", lines[index])
        if bullet:
            document.add_paragraph(re.sub(r"[*`]", "", bullet.group(1)), style="List Bullet")
            index += 1
            continue
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", lines[index])
        if numbered:
            document.add_paragraph(re.sub(r"[*`]", "", numbered.group(1)), style="List Number")
            index += 1
            continue
        if line and not re.fullmatch(r"[-*_]{3,}", line):
            document.add_paragraph(re.sub(r"[*`]", "", line))
        index += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return ToolResult.success(f"Wrote {out_path.name}.", path=str(out_path))


class DocumentTool:
    """`document <request>` — the model writes it, we render it to a real file.

    PDF goes through the same offline Chromium path as the resume export (no network,
    JavaScript disabled); Word needs python-docx and says so plainly when it is absent.
    """

    def __init__(
        self,
        data_dir: Path,
        writer: Writer | None = None,
        approval_gate: ApprovalGate | None = None,
    ) -> None:
        self.directory = Path(data_dir) / "documents"
        self._writer = writer
        self._gate = approval_gate

    def available(self) -> bool:
        return self._writer is not None

    def create(self, spec: str, fmt: str = "pdf") -> ToolResult:
        request, chosen = split_format(spec, fmt)
        if not request:
            return ToolResult.failure(
                "What should the document say? Try 'document a one-page brief on our API rate limits as a pdf'."
            )
        if chosen not in FORMATS:
            return ToolResult.failure(f"I can write pdf, docx or md — not {chosen}.")
        if not self.available():
            return ToolResult.failure("Writing a document needs a language model, and none is configured.")
        if self._gate is not None:  # model call + a file on disk -> MEDIUM, like the image tool
            self._gate.require(
                ApprovalRequest(
                    action=f"Write a {chosen.upper()} document: {request[:120]}",
                    risk=RiskLevel.MEDIUM,
                    reason="Writing a document sends the request to a language model and saves a file locally.",
                )
            )
        try:
            body = (self._writer or (lambda _: ""))(_PROMPT.format(spec=request)) or ""
        except Exception as exc:  # a model outage must not raise out of a tool
            return ToolResult.failure(f"The model could not write that document: {exc}")
        body = body.strip()
        if not body:
            return ToolResult.failure("The model returned an empty document. Try again, or narrow the request.")

        title = _title_of(body, request[:80])
        name = f"{_slug(title)}-{int(time.time())}.{chosen}"
        self.directory.mkdir(parents=True, exist_ok=True)
        out_path = self.directory / name

        if chosen == "md":
            out_path.write_text(body, encoding="utf-8")
            result = ToolResult.success(f"Wrote {name}.", path=str(out_path))
        elif chosen == "docx":
            result = _write_docx(body, out_path)
        else:
            from laptop_agent.tools.resume_pdf import render_html_to_pdf

            html = markdown_to_html(body, title)

            def render() -> ToolResult:
                return asyncio.run(render_html_to_pdf(html, out_path, single_page=False))

            # Check for a running loop before building the coroutine: catching the
            # RuntimeError from asyncio.run() instead would orphan one, which Python
            # reports as "coroutine was never awaited".
            try:
                asyncio.get_running_loop()
            except RuntimeError:
                result = render()
            else:  # inside the web app's loop — render on a worker thread
                import concurrent.futures

                with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
                    result = pool.submit(render).result()
        if not result.ok:
            return result
        url = f"/api/document?name={name}"
        return ToolResult.success(
            f"[{title}]({url}) — {chosen.upper()} ready.\n\n{body[:600]}"
            + ("\n\n…" if len(body) > 600 else ""),
            document=str(out_path),
            name=name,
            url=url,
            title=title,
            format=chosen,
            request=request,
            **({"pages": result.data["pages"]} if "pages" in result.data else {}),
        )
