from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from laptop_agent.safety import ApprovalDenied, ApprovalGate
from laptop_agent.tools.base import ToolResult
from laptop_agent.tools.document import DocumentTool, deck_outline, markdown_to_html, split_format

DECK = """# The Sun and the Planets

## The Sun
- A G2V star holding 99.86% of the system's mass
- Fuses **hydrogen** into helium at about 15 million kelvin

## Mercury
- Smallest planet and closest to the Sun
- One orbit takes 88 Earth days
- No substantial atmosphere
"""

SAMPLE = """# Rate Limit Brief

## Summary
The gateway allows **40** requests per minute.

- Bursts are smoothed over 10s
- A 429 carries `Retry-After`

| Tier | Limit |
|---|---|
| Free | 40/min |
| Paid | 200/min |
"""


class FormatParsingTests(unittest.TestCase):
    def test_a_trailing_format_is_pulled_off_the_request(self) -> None:
        self.assertEqual(split_format("a brief on rate limits as a pdf"), ("a brief on rate limits", "pdf"))
        self.assertEqual(split_format("a brief as a word doc"), ("a brief", "docx"))
        self.assertEqual(split_format("a brief as markdown"), ("a brief", "md"))
        self.assertEqual(split_format("a brief in docx format"), ("a brief", "docx"))

    def test_a_deck_names_its_format_at_the_front(self) -> None:
        """Asked to "create a ppt for sun and planets" the tool wrote a PDF: the format was
        only ever looked for in a tail position ("... as a pdf"), and a deck is not asked
        for that way."""
        self.assertEqual(split_format("a ppt for sun and planets"), ("sun and planets", "pptx"))
        self.assertEqual(split_format("can you create a ppt for sun and planets"), ("sun and planets", "pptx"))
        self.assertEqual(split_format("slides on rust ownership"), ("rust ownership", "pptx"))
        self.assertEqual(split_format("make me a powerpoint on photosynthesis"), ("photosynthesis", "pptx"))
        self.assertEqual(split_format("a deck for Q4 planning"), ("Q4 planning", "pptx"))

    def test_a_deck_named_at_the_tail_still_works(self) -> None:
        self.assertEqual(split_format("the water cycle as a presentation"), ("the water cycle", "pptx"))
        self.assertEqual(split_format("Q4 planning as slides"), ("Q4 planning", "pptx"))

    def test_a_document_request_is_not_mistaken_for_a_deck(self) -> None:
        self.assertEqual(split_format("a report on rust as a pdf"), ("a report on rust", "pdf"))
        self.assertEqual(
            split_format("a one page brief on presentation skills"),
            ("a one page brief on presentation skills", "pdf"),
        )

    def test_no_format_keeps_the_whole_request_and_the_default(self) -> None:
        self.assertEqual(split_format("a brief on rate limits"), ("a brief on rate limits", "pdf"))

    def test_a_format_word_inside_the_topic_is_not_a_format(self) -> None:
        # "how pdf compression works" is the subject, not the output format.
        self.assertEqual(split_format("how pdf compression works"), ("how pdf compression works", "pdf"))


class MarkdownToHtmlTests(unittest.TestCase):
    def test_structure_survives(self) -> None:
        html = markdown_to_html(SAMPLE, "Rate Limit Brief")
        self.assertIn("<h1>Rate Limit Brief</h1>", html)
        self.assertIn("<h2>Summary</h2>", html)
        self.assertIn("<strong>40</strong>", html)
        self.assertIn("<li>Bursts are smoothed over 10s</li>", html)
        self.assertIn("<code>Retry-After</code>", html)
        self.assertIn("<th>Tier</th>", html)
        self.assertIn("<td>Free</td>", html)

    def test_model_output_cannot_inject_markup(self) -> None:
        html = markdown_to_html("# <script>alert(1)</script>\n\nBody <img src=x onerror=1>")
        self.assertNotIn("<script>", html)
        self.assertNotIn("<img", html)
        self.assertIn("&lt;script&gt;", html)

    def test_a_fenced_block_is_kept_verbatim(self) -> None:
        html = markdown_to_html("Run:\n```\nnpm install\n```\nDone.")
        self.assertIn("<pre>npm install</pre>", html)


class DocumentToolTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.data_dir = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def tool(self, writer=lambda prompt: SAMPLE, **kwargs) -> DocumentTool:
        return DocumentTool(data_dir=self.data_dir, writer=writer, **kwargs)

    def test_markdown_is_written_and_linked(self) -> None:
        result = self.tool().create("a brief on rate limits as markdown")
        self.assertTrue(result.ok, result.message)
        path = Path(result.data["document"])
        self.assertTrue(path.is_file())
        self.assertEqual(path.suffix, ".md")
        self.assertEqual(path.read_text(encoding="utf-8"), SAMPLE.strip())
        self.assertEqual(result.data["title"], "Rate Limit Brief")
        self.assertIn(f"]({result.data['url']})", result.message)

    def test_the_request_reaches_the_model_without_the_format_words(self) -> None:
        seen: list[str] = []

        def writer(prompt: str) -> str:
            seen.append(prompt)
            return SAMPLE

        self.tool(writer=writer).create("a brief on rate limits as markdown")
        self.assertIn("a brief on rate limits", seen[0])
        self.assertNotIn("as markdown", seen[0])

    def test_an_empty_request_asks_for_one(self) -> None:
        result = self.tool().create("   as a pdf")
        self.assertFalse(result.ok)
        self.assertIn("What should the document say", result.message)

    def test_an_unsupported_format_is_refused_clearly(self) -> None:
        result = self.tool().create("a brief", fmt="xlsx")
        self.assertFalse(result.ok)
        self.assertIn("pdf, docx, pptx or md", result.message)

    def test_without_a_model_it_says_so(self) -> None:
        result = DocumentTool(data_dir=self.data_dir).create("a brief as markdown")
        self.assertFalse(result.ok)
        self.assertIn("language model", result.message)

    def test_an_empty_model_reply_writes_no_file(self) -> None:
        result = self.tool(writer=lambda prompt: "   ").create("a brief as markdown")
        self.assertFalse(result.ok)
        self.assertIn("empty document", result.message)
        self.assertFalse((self.data_dir / "documents").exists())

    def test_a_model_failure_is_reported_not_raised(self) -> None:
        def broken(prompt: str) -> str:
            raise RuntimeError("model offline")

        result = self.tool(writer=broken).create("a brief as markdown")
        self.assertFalse(result.ok)
        self.assertIn("model offline", result.message)

    def test_documents_written_in_the_same_second_do_not_overwrite_each_other(self) -> None:
        results = [self.tool().create("the weekly brief as markdown") for _ in range(3)]
        names = [r.data["name"] for r in results]
        self.assertEqual(len(set(names)), 3, names)
        self.assertEqual(len(list((self.data_dir / "documents").iterdir())), 3)

    def test_a_failed_render_leaves_no_empty_file_behind(self) -> None:
        # The name is claimed by creating an empty file before rendering, so a render
        # failure must clean it up rather than leaving a 0-byte PDF in the folder.
        from laptop_agent.tools import resume_pdf

        async def broken(html, out_path, single_page=True):
            return ToolResult.failure("Chromium is not installed")

        original = resume_pdf.render_html_to_pdf
        resume_pdf.render_html_to_pdf = broken
        try:
            result = self.tool().create("a brief on rate limits as a pdf")
        finally:
            resume_pdf.render_html_to_pdf = original
        self.assertFalse(result.ok)
        self.assertIn("Chromium", result.message)
        self.assertEqual(list((self.data_dir / "documents").glob("*.pdf")), [])

    def test_denied_approval_stops_the_call(self) -> None:
        calls: list[str] = []

        def writer(prompt: str) -> str:
            calls.append(prompt)
            return SAMPLE

        tool = self.tool(writer=writer, approval_gate=ApprovalGate(ask=lambda request: False))
        with self.assertRaises(ApprovalDenied):
            tool.create("a brief as markdown")
        self.assertEqual(calls, [])


class DeckOutlineTests(unittest.TestCase):
    def test_headings_become_slides_and_bullets_become_lines(self) -> None:
        title, slides = deck_outline(DECK)
        self.assertEqual(title, "The Sun and the Planets")
        self.assertEqual([head for head, _ in slides], ["The Sun", "Mercury"])
        self.assertEqual(len(slides[1][1]), 3)

    def test_emphasis_is_stripped_so_a_slide_shows_words_not_asterisks(self) -> None:
        _title, slides = deck_outline(DECK)
        self.assertIn("Fuses hydrogen into helium at about 15 million kelvin", slides[0][1])

    def test_a_heading_with_no_bullets_is_not_an_empty_slide(self) -> None:
        _title, slides = deck_outline("# T\n\n## Empty\n\n## Real\n- a point\n")
        self.assertEqual([head for head, _ in slides], ["Real"])

    def test_prose_under_a_heading_still_becomes_a_line(self) -> None:
        """The model does sometimes forget the dash. A heading with a paragraph under it is
        a worse slide than one with a prose line on it, but it is not an empty one."""
        _title, slides = deck_outline("# T\n\n## Overview\nThe solar system has eight planets.\n")
        self.assertEqual(slides[0][1], ["The solar system has eight planets."])


class PowerPointTests(unittest.TestCase):
    """Asked for a PPT, the tool produced a PDF — there was no PowerPoint format at all."""

    def setUp(self) -> None:
        self._dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._dir.cleanup)
        self.data_dir = Path(self._dir.name)

    def tool(self, writer=lambda prompt: DECK) -> DocumentTool:
        return DocumentTool(data_dir=self.data_dir, writer=writer)

    def test_the_deck_prompt_asks_for_slides_not_prose(self) -> None:
        seen = []
        self.tool(writer=lambda prompt: (seen.append(prompt), DECK)[1]).create("a ppt for sun and planets")
        self.assertIn("slide deck", seen[0])
        self.assertIn("sun and planets", seen[0])
        self.assertNotIn("Markdown tables", seen[0], "that is the document prompt, not the deck one")

    def test_a_real_pptx_is_written(self) -> None:
        try:
            from pptx import Presentation  # type: ignore
        except ImportError:
            self.skipTest("python-pptx is not installed")
        result = self.tool().create("a ppt for sun and planets")
        self.assertTrue(result.ok, result.message)
        self.assertEqual(result.data["format"], "pptx")
        path = Path(result.data["document"])
        self.assertEqual(path.suffix, ".pptx")

        deck = Presentation(str(path))
        self.assertEqual(len(deck.slides), 3, "a title slide plus one per heading")
        self.assertEqual(deck.slides[0].shapes.title.text, "The Sun and the Planets")
        self.assertEqual(deck.slides[1].shapes.title.text, "The Sun")
        body = [s.text_frame.text for s in deck.slides[2].shapes if s.has_text_frame][1]
        self.assertIn("88 Earth days", body)

    def test_a_deck_with_no_slides_writes_no_file(self) -> None:
        result = self.tool(writer=lambda prompt: "Sorry, I cannot help with that.").create(
            "a ppt for sun and planets"
        )
        self.assertFalse(result.ok)
        self.assertEqual(list(self.data_dir.glob("documents/*.pptx")), [])

    def test_without_python_pptx_it_names_the_library(self) -> None:
        import builtins

        real_import = builtins.__import__

        def no_pptx(name, *args, **kwargs):
            if name == "pptx" or name.startswith("pptx."):
                raise ImportError("No module named 'pptx'")
            return real_import(name, *args, **kwargs)

        builtins.__import__ = no_pptx
        try:
            result = self.tool().create("a ppt for sun and planets")
        finally:
            builtins.__import__ = real_import
        self.assertFalse(result.ok)
        self.assertIn("python-pptx", result.message)
        self.assertIn("pip install", result.message)


if __name__ == "__main__":
    unittest.main()
